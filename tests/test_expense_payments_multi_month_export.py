from __future__ import annotations

import asyncio
import os
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

os.environ.setdefault("SUPABASE_URL", "https://example.supabase.co")
os.environ.setdefault("SUPABASE_KEY", "test-key")
os.environ.setdefault("SUPABASE_SERVICE_ROLE_KEY", "test-service-key")

from routes import gastos_gas_lp


TENANT = "tenant-1"
PROFILE = 7


def _scoped(**values):
    return {"tenant_id": TENANT, "profile_id": PROFILE, **values}


class FakeQuery:
    def __init__(self, rows):
        self.rows = deepcopy(rows)

    def select(self, *_args, **_kwargs):
        return self

    def eq(self, field, value):
        self.rows = [row for row in self.rows if row.get(field) == value]
        return self

    def gte(self, field, value):
        self.rows = [row for row in self.rows if str(row.get(field) or "") >= str(value)]
        return self

    def lt(self, field, value):
        self.rows = [row for row in self.rows if str(row.get(field) or "") < str(value)]
        return self

    def in_(self, field, values):
        allowed = set(values)
        self.rows = [row for row in self.rows if row.get(field) in allowed]
        return self

    def order(self, field, desc=False):
        self.rows.sort(key=lambda row: row.get(field) or "", reverse=desc)
        return self

    def limit(self, count):
        self.rows = self.rows[:count]
        return self

    def execute(self):
        return SimpleNamespace(data=deepcopy(self.rows))


class FakeSupabase:
    def __init__(self, tables):
        self.tables = tables

    def table(self, name):
        return FakeQuery(self.tables.get(name, []))


def _xlsx_bytes(response):
    async def collect():
        chunks = []
        async for chunk in response.body_iterator:
            chunks.append(chunk)
        return b"".join(chunks)

    return asyncio.run(collect())


def test_accounting_export_includes_every_requested_payment_type_and_excludes_unselected_month(monkeypatch):
    tables = {
        "gas_lp_expense_payments": [
            _scoped(id=1, paid_on="2026-06-10", amount_mxn=1500, payment_target="supplier", supplier_id=10,
                    reimbursement_recipient_id=None, reimbursement_account_id=None, reference="JUN-001", notes="Pago junio"),
            _scoped(id=2, paid_on="2026-07-10", amount_mxn=999, payment_target="supplier", supplier_id=10,
                    reimbursement_recipient_id=None, reimbursement_account_id=None, reference="JUL-EXCLUIR", notes=""),
            _scoped(id=3, paid_on="2026-08-20", amount_mxn=500, payment_target="reimbursement", supplier_id=None,
                    reimbursement_recipient_id=20, reimbursement_account_id=30, reference="AGO-001", notes="Taxi"),
        ],
        "gas_lp_expense_advances": [
            _scoped(id=200, paid_on="2026-09-05", amount_mxn=700, supplier_id=10, concept_id=2,
                    status="pending", reference="SEP-ANT", description="Anticipo septiembre"),
        ],
        "gas_lp_expense_advance_applications": [],
        "gas_lp_expense_payment_allocations": [
            {"payment_id": 1, "invoice_id": 100, "amount_mxn": 1000},
            {"payment_id": 1, "invoice_id": 101, "amount_mxn": 500},
            {"payment_id": 2, "invoice_id": 102, "amount_mxn": 999},
            {"payment_id": 3, "invoice_id": 103, "amount_mxn": 500},
        ],
        "gas_lp_expense_invoices": [
            _scoped(id=100, supplier_id=10, concept_id=1, invoice_number="JUN-A", invoice_date="2026-06-01", total_mxn=1000),
            _scoped(id=101, supplier_id=10, concept_id=2, invoice_number="JUN-B", invoice_date="2026-06-02", total_mxn=600),
            _scoped(id=102, supplier_id=10, concept_id=1, invoice_number="JUL-NO", invoice_date="2026-07-01", total_mxn=999),
            _scoped(id=103, supplier_id=11, concept_id=2, invoice_number="AGO-R", invoice_date="2026-08-15", total_mxn=500),
        ],
        "gas_lp_expense_suppliers": [
            _scoped(id=10, commercial_name="Proveedor Uno", legal_name="Proveedor Uno SA", bank_name="BBVA", account_number="1234"),
            _scoped(id=11, commercial_name="Taxi Centro", legal_name="Taxi Centro SA", bank_name="", account_number=""),
        ],
        "gas_lp_expense_recipients": [
            _scoped(id=20, name="Ana Pérez", email="ana@example.com"),
        ],
        "gas_lp_expense_recipient_accounts": [
            _scoped(id=30, recipient_id=20, account_type="credit_card", bank_name="Banorte", card_last_four="4455"),
        ],
        "gas_lp_expense_concepts": [
            _scoped(id=1, name="Combustible"),
            _scoped(id=2, name="Traslados"),
        ],
    }
    fake_sb = FakeSupabase(tables)
    context = {
        "sb": fake_sb,
        "tenant_id": TENANT,
        "perfil_id": PROFILE,
        "actor_id": "user-1",
        "is_admin": True,
    }
    monkeypatch.setattr(gastos_gas_lp, "_ctx", lambda *_args, **_kwargs: context)
    monkeypatch.setattr(gastos_gas_lp, "_profile", lambda _ctx: {"nombre": "Empresa de prueba"})

    response = gastos_gas_lp.export_expense_payments(
        token="", month="", months="2026-06,2026-08,2026-09",
        authorization="", x_flotilla_access="", x_perfil_id=str(PROFILE),
    )
    workbook = load_workbook(BytesIO(_xlsx_bytes(response)), data_only=False)

    assert workbook.sheetnames == ["Pagos proveedores", "Reembolsos", "Resumen por destinatario"]
    provider = workbook["Pagos proveedores"]
    reimbursement = workbook["Reembolsos"]
    summary = workbook["Resumen por destinatario"]

    expected_headers = [
        "ID DE PAGO", "FACTURA EN EL PAGO", "FOLIO DE FACTURA", "FECHA DE FACTURA",
        "CONCEPTO DE GASTO", "PROVEEDOR / EMISOR", "REEMBOLSADO A", "DESTINO DEL PAGO",
        "MONTO DE LA FACTURA", "MONTO PAGADO", "FECHA DE PAGO", "TOTAL TRANSFERIDO",
        "DIFERENCIA", "REFERENCIA", "NOTAS",
    ]
    assert [cell.value for cell in provider[1]] == expected_headers
    assert [cell.value for cell in reimbursement[1]] == expected_headers

    provider_rows = list(provider.iter_rows(min_row=2, values_only=True))
    reimbursement_rows = list(reimbursement.iter_rows(min_row=2, values_only=True))
    assert [row[2] for row in provider_rows] == ["JUN-A", "JUN-B", "S/F"]
    assert [row[0] for row in provider_rows] == ["P-1", "P-1", "A-200"]
    assert [row[2] for row in reimbursement_rows] == ["AGO-R"]
    assert reimbursement_rows[0][6] == "Ana Pérez"
    assert reimbursement_rows[0][7] == "Tarjeta •••• 4455"
    assert "PAGO PARCIAL" in provider_rows[1][14]
    assert "PENDIENTE DE FACTURA" in provider_rows[2][14]
    assert all("JUL" not in str(value) for row in provider_rows + reimbursement_rows for value in row)

    exported_paid = sum(float(row[9] or 0) for row in provider_rows + reimbursement_rows)
    assert exported_paid == 2700
    summary_rows = list(summary.iter_rows(min_row=2, max_row=summary.max_row - 1, values_only=True))
    assert sum(float(row[4] or 0) for row in summary_rows) == 2700
    assert summary.cell(summary.max_row, 2).value == "TOTAL DEL PERIODO"
    assert response.headers["content-disposition"] == (
        'attachment; filename="relacion-pagos-2026-06_a_2026-09-3-meses.xlsx"'
    )

